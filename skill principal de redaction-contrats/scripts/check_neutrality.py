#!/usr/bin/env python3
"""
check_neutrality.py — Détecte les résidus d'institutions, de noms propres ou de
vocabulaire propre au secteur public dans un contrat (.md ou .docx).

Usage :
    python check_neutrality.py contrat.docx [--extra "Nom Société" --extra "Autre"] [--public]

--public : autorise le vocabulaire secteur public (quand l'utilisateur a demandé
           explicitement la variante administration). Les noms propres restent interdits.
Code de sortie : 0 = propre, 1 = résidus trouvés.
"""
import re
import sys
import argparse

# Noms propres / structures issus des modèles sources : jamais dans un contrat neutre.
ALWAYS = [
    r"Contr[ôo]le Financier", r"N.?DA\s*KACOU", r"Joseph\s*Anoma", r"SMGL",
    r"\bDCF\b", r"\bPPRC\b", r"\bC2D\b", r"\bCOCAN\b", r"\bPRIME\b", r"\bPDU\b",
    r"\bDGBF\b", r"\bDMP\b", r"Armoiries", r"Minist[èe]re", r"01\s*BP\s*V\s*80",
    r"20\s*21\s*51\s*40", r"20\s*21\s*34\s*87", r"PROPOSITION DE CONTENU",
    r"\bCNPS\b", r"Caisse Nationale de Pr[ée]voyance",
]
# Vocabulaire administratif : à éviter dans un modèle neutre (autorisé avec --public).
PUBLIC = [
    r"Autorit[ée] contractante", r"budget de l.[ÉE]tat", r"imputation budg[ée]taire",
    r"ligne budg[ée]taire", r"Fonction Publique", r"service public",
    r"int[ée]r[êe]t g[ée]n[ée]ral", r"\bdestination\b\s*[….]",
]


def read_text(path):
    if path.lower().endswith(".docx"):
        from docx import Document
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts += [c.text for c in row.cells]
        for s in d.sections:
            parts += [p.text for p in s.header.paragraphs + s.footer.paragraphs]
        return "\n".join(parts)
    with open(path, encoding="utf-8") as f:
        return f.read()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("fichier")
    ap.add_argument("--extra", action="append", default=[],
                    help="terme supplémentaire interdit (ex. nom d'une société citée dans les sources)")
    ap.add_argument("--public", action="store_true")
    a = ap.parse_args()

    text = read_text(a.fichier)
    patterns = ALWAYS + [re.escape(e) for e in a.extra] + ([] if a.public else PUBLIC)
    hits = []
    for pat in patterns:
        for m in re.finditer(pat, text, flags=re.I if pat in PUBLIC else 0):
            s = max(0, m.start() - 40)
            hits.append((m.group(0), text[s:m.end() + 40].replace("\n", " ")))
    if hits:
        print(f"{len(hits)} résidu(s) non neutre(s) dans {a.fichier} :")
        for term, ctx in hits:
            print(f"  - « {term} » …{ctx}…")
        sys.exit(1)
    print(f"Neutralité OK : {a.fichier}")


if __name__ == "__main__":
    main()
