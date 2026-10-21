#!/usr/bin/env python3
"""
build_contract.py — Convertit un contrat rédigé en "Markdown contrat" (voir
references/format-source.md) en fichier .docx mis en forme.

Usage :
    python build_contract.py contrat.md sortie.docx [--format A4|Letter] [--sans-sommaire]

Mise en forme : Arial 10 pt, texte justifié, articles numérotés automatiquement,
sommaire statique, pied de page « Page X sur Y », tableaux, bloc de signatures.
Les champs à compléter ([entre crochets] ou points de suspension) sont surlignés
en jaune pour qu'ils ne passent pas inaperçus avant signature.
"""
import re
import argparse
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Arial"
PLACEHOLDER = re.compile(r"(\[[^\]\n]*\]|…{2,}|\.{4,})")
BOLD = re.compile(r"(\*\*.+?\*\*)")


# ---------------------------------------------------------------- helpers ---
def set_font(run, size=None, bold=None, italic=None, underline=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline


def add_inline(par, text, size=10, bold=False, italic=False, underline=False):
    """Ajoute du texte avec **gras** et surlignage des champs à compléter."""
    for chunk in BOLD.split(text):
        if not chunk:
            continue
        is_b = bold
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            chunk, is_b = chunk[2:-2], True
        for piece in PLACEHOLDER.split(chunk):
            if not piece:
                continue
            r = par.add_run(piece)
            set_font(r, size, is_b, italic, underline)
            if PLACEHOLDER.fullmatch(piece):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def add_field(par, instr, size=8, italic=False):
    """Insère un champ Word (PAGE, NUMPAGES) sous forme de champ complexe."""
    def new_run():
        r = par.add_run()
        set_font(r, size, italic=italic)
        return r

    r = new_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r._element.append(fc)

    r = new_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._element.append(it)

    r = new_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    r._element.append(fc)

    r = new_run()
    t = OxmlElement("w:t")
    t.text = "1"
    r._element.append(t)

    r = new_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r._element.append(fc)


def para(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=10, bold=False,
         italic=False, underline=False, before=0, after=6, keep_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.15
    pf.keep_with_next = keep_next
    if text:
        add_inline(p, text, size, bold, italic, underline)
    return p


def set_cell_borders(cell, val="single", sz=4):
    tcPr = cell._element.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), "808080" if val != "nil" else "auto")
        b.append(e)
    tcPr.append(b)


def shade(cell, fill="E7E6E6"):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def add_table(doc, rows, content_width_cm):
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    # pondération des colonnes : « N° » étroite, 2e colonne élargie sur les grands tableaux
    weights = []
    for j in range(ncols):
        h = rows[0][j] if j < len(rows[0]) else ""
        wgt = 0.75 if h.strip().lower() in ("n°", "n", "no") else 1.0
        if j == 1 and ncols >= 6:
            wgt = 2.4
        weights.append(wgt)
    tot = sum(weights)
    widths = [Cm(content_width_cm * w / tot) for w in weights]
    size = 8 if ncols >= 8 else 9
    for j, col in enumerate(t.columns):
        col.width = widths[j]
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = t.cell(i, j)
            cell.width = widths[j]
            set_cell_borders(cell)
            txt = row[j] if j < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            add_inline(p, txt, size, bold=(i == 0))
            if i == 0:
                shade(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_signatures(doc, labels, content_width_cm):
    t = doc.add_table(rows=1, cols=len(labels))
    t.autofit = False
    w = Cm(content_width_cm / len(labels))
    for j, lab in enumerate(labels):
        c = t.cell(0, j)
        c.width = w
        set_cell_borders(c, "nil")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, lab.strip(), 10, bold=True)
        p2 = c.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p2, "(Nom, qualité, signature précédée de la mention « Lu et approuvé »)", 8, italic=True)
        for _ in range(5):
            c.add_paragraph()


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ----------------------------------------------------------------- parsing ---
def parse_source(text):
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    meta, body = {}, text
    m = re.match(r"\s*---\n(.*?)\n---\n(.*)", text, flags=re.S)
    if m:
        for line in m.group(1).splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip().lower()] = v.strip()
        body = m.group(2)
    return meta, body.splitlines()


def build(src_path, out_path, fmt="A4", sommaire=None):
    with open(src_path, encoding="utf-8") as f:
        meta, lines = parse_source(f.read())

    doc = Document()
    sec = doc.sections[0]
    if (meta.get("format") or fmt).lower() == "letter":
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    else:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    content_w = (sec.page_width - sec.left_margin - sec.right_margin) / 360000  # EMU → cm

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    h1 = doc.styles["Heading 1"]
    h1.font.name, h1.font.size, h1.font.bold = FONT, Pt(10.5), True
    h1.font.color.rgb = RGBColor(0, 0, 0)

    # pied de page : « Page X sur Y »
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    set_font(r, 8)
    add_field(fp, "PAGE")
    r = fp.add_run(" sur ")
    set_font(r, 8)
    add_field(fp, "NUMPAGES")

    # 1ère passe : liste des articles pour le sommaire
    titles = [l[3:].strip() for l in lines if l.startswith("## ")]
    want_toc = (meta.get("sommaire", "oui").lower() in ("oui", "true", "1")) and sommaire is not False

    state = {"title": False, "toc": False, "art": 0}
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in table_buf
                    if not re.fullmatch(r"\|?[\s:\-|]+\|?", r)]
            if rows:
                add_table(doc, rows, content_w)
            table_buf = []

    def emit_toc():
        if state["toc"] or not (want_toc and titles):
            return
        state["toc"] = True
        para(doc, "SOMMAIRE", WD_ALIGN_PARAGRAPH.LEFT, 10, True, before=6, after=6)
        for i, t in enumerate(titles, 1):
            para(doc, f"ARTICLE {i} : {t.upper()}", WD_ALIGN_PARAGRAPH.LEFT, 9, after=2)
        page_break(doc)

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("|"):
            table_buf.append(line)
            continue
        flush_table()
        if not line.strip():
            continue

        if line.startswith("# ") and not state["title"]:
            state["title"] = True
            para(doc, line[2:].strip().upper(), WD_ALIGN_PARAGRAPH.CENTER, 14, True, before=12, after=4)
            if meta.get("sous-titre"):
                para(doc, meta["sous-titre"], WD_ALIGN_PARAGRAPH.CENTER, 10, italic=True, after=14)
            emit_toc()
        elif line.startswith("## "):
            emit_toc()
            state["art"] += 1
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(f"ARTICLE {state['art']} : {line[3:].strip().upper()}")
            set_font(r, 10.5, True)
            r.font.color.rgb = RGBColor(0, 0, 0)
        elif line.startswith("### "):
            para(doc, line[4:].strip(), WD_ALIGN_PARAGRAPH.LEFT, 10, True, before=4, after=4, keep_next=True)
        elif line.startswith("@annexe "):
            page_break(doc)
            para(doc, line[8:].strip(), WD_ALIGN_PARAGRAPH.LEFT, 11, True, underline=True, after=8, keep_next=True)
        elif line.startswith("@fait "):
            para(doc, line[6:].strip(), WD_ALIGN_PARAGRAPH.LEFT, 10, before=14, after=8, keep_next=True)
        elif line.startswith("@signatures "):
            add_signatures(doc, line[12:].split("|"), content_w)
        elif line.startswith("@paraphe"):
            p = para(doc, "", WD_ALIGN_PARAGRAPH.CENTER, 8, italic=True, before=8)
            r = p.add_run("(Ce contrat comporte ")
            set_font(r, 8, italic=True)
            add_field(p, "NUMPAGES", italic=True)
            r = p.add_run(" pages paraphées par les parties)")
            set_font(r, 8, italic=True)
        elif line.startswith("@saut"):
            page_break(doc)
        elif line.startswith("@centre "):
            para(doc, line[8:].strip(), WD_ALIGN_PARAGRAPH.CENTER, 10, True)
        elif line.startswith("~"):  # étiquettes : Entre / D'une part / Et / D'autre part
            para(doc, line[1:].strip(), WD_ALIGN_PARAGRAPH.LEFT, 10, True, before=6, after=6, keep_next=True)
        elif re.match(r"^\s*[-*] ", line):
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, re.sub(r"^\s*[-*] ", "", line))
        else:
            para(doc, line.strip())
    flush_table()
    emit_toc()
    # le gabarit python-docx déclare <w:zoom/> sans pourcentage : on le renseigne (validité OOXML)
    for z in doc.settings.element.findall(qn("w:zoom")):
        if z.get(qn("w:percent")) is None:
            z.set(qn("w:percent"), "100")
    doc.save(out_path)
    print(f"OK -> {out_path}  ({state['art']} articles)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Markdown contrat -> .docx")
    ap.add_argument("source")
    ap.add_argument("sortie")
    ap.add_argument("--format", default="A4", choices=["A4", "Letter"])
    ap.add_argument("--sans-sommaire", action="store_true")
    a = ap.parse_args()
    build(a.source, a.sortie, a.format, False if a.sans_sommaire else None)
